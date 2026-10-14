"""
Generate Synoptic Presentation PDF and PPTX from User-Provided Slides
Author: Azhar Fareed Mulla (USN: 2SA25MC002)
Institution: SECAB Institute of Engineering & Technology, Vijayapur
"""
import os
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def generate_pdf(pdf_path):
    # Landscape orientation for slides format
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=landscape(letter),
        rightMargin=45, leftMargin=45,
        topMargin=35, bottomMargin=35
    )

    styles = getSampleStyleSheet()

    title_inst = ParagraphStyle(
        'InstTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=19,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#8B0000')
    )

    mca_sub = ParagraphStyle(
        'MCASub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#1A5276')
    )

    proj_title = ParagraphStyle(
        'ProjTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=25,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#900C3F')
    )

    meta_info = ParagraphStyle(
        'MetaInfo',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=17,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#2C3E50')
    )

    slide_num = ParagraphStyle(
        'SlideNum',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#8B0000')
    )

    slide_header = ParagraphStyle(
        'SlideHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1C2833')
    )

    bullet_text = ParagraphStyle(
        'BulletText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=18,
        leftIndent=20,
        textColor=colors.HexColor('#212F3D')
    )

    sub_bullet_text = ParagraphStyle(
        'SubBulletText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=16,
        leftIndent=40,
        textColor=colors.HexColor('#4A235A')
    )

    box_text = ParagraphStyle(
        'BoxText',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1E8449')
    )

    story = []

    # ── SLIDE 1: Title ──
    story.append(Spacer(1, 40))
    story.append(Paragraph("SECAB Institute of Engineering & Technology, Vijayapur", title_inst))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Department of Master of Computer Applications (MCA) | Synoptic Note 2026–27", mca_sub))
    story.append(Spacer(1, 25))
    story.append(HRFlowable(width="80%", thickness=2, color=colors.HexColor('#900C3F'), spaceAfter=25))
    story.append(Paragraph("Intelligent E-Waste Classification and Recovery Recommendation Using Data Science", proj_title))
    story.append(Spacer(1, 30))
    story.append(Paragraph("<b>Student:</b> Azhar Fareed Mulla (USN: <b>2SA25MC002</b>) &nbsp;&nbsp;|&nbsp;&nbsp; <b>Semester:</b> III", meta_info))
    story.append(Spacer(1, 8))
    story.append(Paragraph("<b>Date:</b> 04/09/26 &nbsp;&nbsp;|&nbsp;&nbsp; <b>Guide:</b> Dr. Nisha S. Amin", meta_info))
    story.append(PageBreak())

    # ── SLIDE 2: Abstract ──
    story.append(Paragraph("Slide 2", slide_num))
    story.append(Paragraph("Abstract", slide_header))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#BDC3C7'), spaceAfter=20))
    story.append(Paragraph("• E-waste is a fast-growing waste stream containing recoverable materials and hazardous substances", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Many centers rely on manual sorting and generic disposal rules", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Proposed system uses machine learning to:", bullet_text))
    story.append(Paragraph("— Classify e-waste categories accurately", sub_bullet_text))
    story.append(Paragraph("— Estimate hazard levels for safe handling", sub_bullet_text))
    story.append(Paragraph("— Identify likely recoverable components (precious metals & parts)", sub_bullet_text))
    story.append(Paragraph("— Recommend the best recycling/recovery pathway", sub_bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Multiple models will be trained and compared; feature engineering will be used for meaningful predictors", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Deliverable includes trained models, performance comparison, feature influence analysis, and a lightweight reporting interface", bullet_text))
    story.append(PageBreak())

    # ── SLIDE 3: Introduction ──
    story.append(Paragraph("Slide 3", slide_num))
    story.append(Paragraph("Introduction", slide_header))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#BDC3C7'), spaceAfter=20))
    story.append(Paragraph("• Global e-waste volume is rising due to frequent replacement of electronics and IT hardware", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Different items require distinct, specialized handling:", bullet_text))
    story.append(Paragraph("— Recoverable precious metals (Gold, Silver, Copper, Platinum, Palladium)", sub_bullet_text))
    story.append(Paragraph("— Hazardous components (e.g., Lithium-ion batteries, CRT glass, heavy metals)", sub_bullet_text))
    story.append(Paragraph("— Structural plastics, aluminum, and low-value recyclable materials", sub_bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Current intake process depends heavily on manual inspection and operator experience", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• This manual approach is slow, inconsistent, and difficult to scale across industrial volumes", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Data science and machine learning can predict category and hazard level using item attributes before full teardown", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• The system explains which attributes influenced the prediction to build operator trust and usability", bullet_text))
    story.append(PageBreak())

    # ── SLIDE 4: Background and Motivation ──
    story.append(Paragraph("Slide 4", slide_num))
    story.append(Paragraph("Background and Motivation", slide_header))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#BDC3C7'), spaceAfter=18))
    story.append(Paragraph("• Many facilities use manual sorting or rigid fixed rules (e.g., treating all batteries uniformly)", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• These approaches struggle with the sheer diversity and incoming volume of electronic hardware", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Desired System Capabilities:</b>", bullet_text))
    story.append(Paragraph("— Automated classification into standard e-waste categories", sub_bullet_text))
    story.append(Paragraph("— Leverage item attributes: device type, weight, material composition, vintage, and condition", sub_bullet_text))
    story.append(Paragraph("— Comparison of multiple machine learning models for superior accuracy", sub_bullet_text))
    story.append(Paragraph("— Hazard level estimation for safe, injury-free operational handling", sub_bullet_text))
    story.append(Paragraph("— Identification of recoverable precious components & recommended custom recovery process", sub_bullet_text))
    story.append(Paragraph("— Feature importance insight driving transparent decision-making across shifts", sub_bullet_text))
    story.append(Spacer(1, 10))
    
    # Example Callout Box
    ex_table = Table([[
        Paragraph("<b>Example Output:</b><br/>Lithium-ion Battery &rarr; <b>Category:</b> Battery Waste &nbsp;|&nbsp; <b>Hazard:</b> High &nbsp;|&nbsp; <b>Recoverable:</b> Lithium, Cobalt, Copper &nbsp;|&nbsp; <b>Recommended:</b> Specialized Closed-Loop Battery Recycling", box_text)
    ]], colWidths=[680])
    ex_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EAFAF1')),
        ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor('#27AE60')),
        ('PADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(ex_table)
    story.append(PageBreak())

    # ── SLIDE 5: Problem Statement ──
    story.append(Paragraph("Slide 5", slide_num))
    story.append(Paragraph("Problem Statement", slide_header))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#BDC3C7'), spaceAfter=20))
    story.append(Paragraph("• E-waste intake is rarely classified or triaged in a systematic, data-driven framework", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• Existing sorting is largely manual and does not consistently identify:", bullet_text))
    story.append(Paragraph("— Accurate hazard exposure levels", sub_bullet_text))
    story.append(Paragraph("— Embedded high-value recoverable components", sub_bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• This slows facility processing times and increases risks of unsafe handling or informal landfill dumping", bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• <b>Critical Need:</b> An integrated system that:", bullet_text))
    story.append(Paragraph("— Classifies incoming items early using machine learning on recorded attributes", sub_bullet_text))
    story.append(Paragraph("— Reliably estimates hazard levels before physical exposure", sub_bullet_text))
    story.append(Paragraph("— Maps out recoverable components and advises optimal recovery pathways", sub_bullet_text))
    story.append(Spacer(1, 8))
    story.append(Paragraph("• All outputs must be mathematically trusted, transparent, and actionable for recycling operators", bullet_text))
    story.append(PageBreak())

    # ── SLIDE 6: Need for the Project ──
    story.append(Paragraph("Slide 6", slide_num))
    story.append(Paragraph("Need for the Project", slide_header))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#BDC3C7'), spaceAfter=20))
    story.append(Paragraph("• <b>Speed & Scalability:</b> Manual sorting is slow, inconsistent, and cannot keep up with growing intake volumes", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Workplace Safety:</b> Hazard levels vary widely; misclassification creates critical fire, chemical, and environmental hazards", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Leverage Existing Intake Records:</b> Facilities already track basic physical attributes that can fuel automated AI", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Inter-feature Intelligence:</b> Generic rules ignore complex interactions between item weight, age, and condition", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Objective Decision Making:</b> Data-driven algorithms reduce reliance on subjective, informal operator judgment", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Resource Maximization:</b> Accurate component identification dramatically improves secondary material yield", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Multi-Model Comparison:</b> Validating multiple ML models ensures selection of the most reliable predictor", bullet_text))
    story.append(Spacer(1, 6))
    story.append(Paragraph("• <b>Shift Standardization:</b> Establishes uniform operational accuracy across different shifts and recycling facilities", bullet_text))

    doc.build(story)

def generate_pptx(pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    c_red = RGBColor(144, 12, 63)
    c_dark = RGBColor(28, 40, 51)
    c_blue = RGBColor(26, 82, 118)
    c_gray = RGBColor(44, 62, 80)
    c_sub = RGBColor(74, 35, 90)

    # ── Slide 1 ──
    s1 = prs.slides.add_slide(blank_layout)
    tx_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SECAB Institute of Engineering & Technology, Vijayapur"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = c_red
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Department of Master of Computer Applications (MCA) | Synoptic Note 2026–27"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = c_blue
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = "\nIntelligent E-Waste Classification and Recovery Recommendation Using Data Science"
    p3.font.size = Pt(28)
    p3.font.bold = True
    p3.font.color.rgb = c_dark
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf.add_paragraph()
    p4.text = "\nStudent: Azhar Fareed Mulla (USN: 2SA25MC002)   |   Semester: III\nDate: 04/09/26   |   Guide: Dr. Nisha S. Amin"
    p4.font.size = Pt(15)
    p4.font.color.rgb = c_gray
    p4.alignment = PP_ALIGN.CENTER

    def add_content_slide(title, points):
        s = prs.slides.add_slide(blank_layout)
        t_box = s.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        t_frame = t_box.text_frame
        p_t = t_frame.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = c_red

        c_box = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.4))
        c_frame = c_box.text_frame
        c_frame.word_wrap = True

        for idx, pt in enumerate(points):
            p = c_frame.paragraphs[0] if idx == 0 else c_frame.add_paragraph()
            if pt.startswith("  "):
                p.text = "    —  " + pt.strip()
                p.font.size = Pt(17)
                p.font.color.rgb = c_sub
            else:
                p.text = "•  " + pt.strip()
                p.font.size = Pt(19)
                p.font.color.rgb = c_dark
            p.space_after = Pt(8)

    # ── Slide 2 ──
    add_content_slide("Abstract", [
        "E-waste is a fast-growing waste stream containing recoverable materials and hazardous substances",
        "Many centers rely on manual sorting and generic disposal rules",
        "Proposed system uses machine learning to:",
        "  classify e-waste categories",
        "  estimate hazard level",
        "  identify likely recoverable components",
        "  recommend the best recycling/recovery pathway",
        "Multiple models will be trained and compared; feature engineering will be used for meaningful predictors",
        "Deliverable includes trained models, performance comparison, feature influence analysis, and a lightweight reporting interface"
    ])

    # ── Slide 3 ──
    add_content_slide("Introduction", [
        "Global e-waste volume is rising due to frequent replacement of electronics and IT hardware",
        "Different items require different handling:",
        "  recoverable precious metals",
        "  hazardous components (e.g., batteries, CRT glass)",
        "  plastics and low-value materials",
        "Current process depends on manual inspection and operator experience",
        "This approach is slow, inconsistent, and difficult to scale",
        "Data science and machine learning can predict category and hazard level using item attributes before full inspection",
        "The system should also explain which attributes influenced the prediction to build trust and usability"
    ])

    # ── Slide 4 ──
    add_content_slide("Background and Motivation", [
        "Many facilities use manual sorting or fixed rules (e.g., treating all batteries as hazardous)",
        "These approaches struggle with the diversity and volume of incoming items",
        "Desired system capabilities:",
        "  automated classification into standard e-waste categories",
        "  use item attributes such as type, weight, material composition, and age",
        "  comparison of multiple machine learning models for best accuracy",
        "  hazard level estimation for safe handling",
        "  identification of likely recoverable components",
        "  recommended recovery process tailored to the item",
        "  insight into features driving decisions (transparency)",
        "  repeatable process for consistent results across large intake volumes",
        "Example output: Lithium-ion Battery → Battery Waste | Hazard: High | Recoverable: Lithium, Cobalt, Copper | Recommended: Specialized Battery Recycling"
    ])

    # ── Slide 5 ──
    add_content_slide("Problem Statement", [
        "E-waste intake is rarely classified or triaged in a systematic, data-driven way",
        "Existing sorting is largely manual and does not consistently identify:",
        "  hazard levels",
        "  recoverable components",
        "This slows processing and increases risk of unsafe or inefficient handling",
        "Need a system that:",
        "  classifies incoming items early using machine learning on attributes",
        "  estimates hazard level",
        "  identifies recoverable components",
        "  recommends an appropriate recovery process",
        "The outputs must be trusted and actionable for recycling operators"
    ])

    # ── Slide 6 ──
    add_content_slide("Need for the Project", [
        "Manual sorting is slow, inconsistent, and difficult to scale",
        "Hazard levels vary widely; misclassification can create safety or environmental risks",
        "Facilities already record basic item attributes that can be leveraged for automation",
        "Generic one-size-fits-all rules do not account for interactions between item attributes",
        "Data-driven recommendations reduce reliance on informal operator judgment",
        "Recoverable component identification improves material recovery and resource efficiency",
        "Comparing multiple models helps select the most reliable approach for a facility’s data",
        "A standardized process improves consistency across different shifts and operators"
    ])

    prs.save(pptx_path)

if __name__ == '__main__':
    pdf_out = r"C:\Users\HP\Desktop\Synoptic_Note_Presentation.pdf"
    pptx_out = r"C:\Users\HP\Desktop\Synoptic_Note_Presentation.pptx"
    
    generate_pdf(pdf_out)
    generate_pptx(pptx_out)
    print("GENERATED PDF:", pdf_out)
    print("GENERATED PPTX:", pptx_out)
